"""Build the Excel workbook and the site data module from figures-data.json.

figures-data.json is the single source of truth for every figure transcribed
from the report. Edit values there, then rerun:

    python data-extraction/build.py

Outputs:
    data-extraction/ipr25-figures.xlsx  - index + one sheet per figure + long-format data + flags
    src/data/figures.generated.js       - same data in the shape statistics.js reads

Needs openpyxl (pip install openpyxl).
"""

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
DATA = HERE / "figures-data.json"
XLSX = HERE / "ipr25-figures.xlsx"
JS = ROOT / "src" / "data" / "figures.generated.js"

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)
HEADER_FILL = PatternFill("solid", fgColor="E6F2F0")
FLAG_FONT = Font(color="B3261E")
WRAP = Alignment(wrap_text=True, vertical="top")
HEX = re.compile(r"#[0-9a-f]{6}")


def bodies(fig):
    """(title, body) for each chart or table in the figure: its panels, or the figure itself."""
    if "panels" in fig:
        return [(panel.get("title"), panel) for panel in fig["panels"]]
    return [(None, fig)]


def body_kind(body):
    if "table" in body:
        return "table"
    chart = body["chart"]
    kind = chart["type"]
    if chart.get("horizontal"):
        kind = "horizontal " + kind
    if chart.get("stacked"):
        kind = "stacked " + kind
    return kind


def figure_kind(fig):
    if "skipped" in fig:
        return "skipped"
    kinds = [body_kind(body) for _, body in bodies(fig)]
    if len(kinds) == 1:
        return kinds[0]
    return "panels: " + " + ".join(kinds)


def chart_problems(where, chart):
    problems = []
    n = len(chart["labels"])
    if "values" in chart and len(chart["values"]) != n:
        problems.append(f"{where}: {len(chart['values'])} values for {n} labels")
    for ds in chart.get("datasets", []):
        if len(ds["data"]) != n:
            problems.append(f"{where}: dataset '{ds['label']}' has {len(ds['data'])} points for {n} labels")
        # color is one colour for the series, or one per data point.
        color = ds.get("color")
        if isinstance(color, list) and len(color) != n:
            problems.append(f"{where}: dataset '{ds['label']}' has {len(color)} colours for {n} labels")
        for c in color if isinstance(color, list) else [color] if color else []:
            if not HEX.fullmatch(c):
                problems.append(f"{where}: dataset '{ds['label']}' has bad colour {c!r}")
    if "colors" in chart:
        if len(chart["colors"]) != n:
            problems.append(f"{where}: {len(chart['colors'])} colours for {n} labels")
        problems += [f"{where}: bad colour {c!r}" for c in chart["colors"] if not HEX.fullmatch(c)]
    return problems


def table_problems(where, table):
    n = len(table["columns"])
    return [
        f"{where}: row {row[0]!r} has {len(row)} cells for {n} columns"
        for row in table["rows"] + ([table["total"]] if "total" in table else [])
        if len(row) != n
    ]


def validate(figures):
    """Fail loudly on shape mismatches so a typo can't reach the site."""
    seen = set()
    problems = []
    for fig in figures:
        fid = fig["id"]
        if fid in seen:
            problems.append(f"{fid}: duplicate id")
        seen.add(fid)
        if "skipped" not in fig and sum(k in fig for k in ("chart", "table", "panels")) != 1:
            problems.append(f"{fid}: needs exactly one of chart, table or panels")
        for title, body in bodies(fig):
            where = f"{fid} [{title}]" if title else fid
            if "chart" in body:
                problems += chart_problems(where, body["chart"])
            elif "table" in body:
                problems += table_problems(where, body["table"])
            elif "skipped" not in fig:
                problems.append(f"{where}: no chart or table")
    if problems:
        raise SystemExit("figures-data.json has problems:\n  " + "\n  ".join(problems))


def body_rows(body):
    """A chart's or table's data as a header row plus body rows (and an optional total row)."""
    if "table" in body:
        table = body["table"]
        return [c["label"] for c in table["columns"]], table["rows"], table.get("total")
    chart = body["chart"]
    if "values" in chart:
        return ["Category", "Value"], [[l, v] for l, v in zip(chart["labels"], chart["values"])], None
    header = ["Category"] + [ds["label"] for ds in chart["datasets"]]
    rows = [[label] + [ds["data"][i] for ds in chart["datasets"]] for i, label in enumerate(chart["labels"])]
    return header, rows, None


def body_long_rows(body, default_series):
    """One row per value: (series, category, value). Tables use column label as series."""
    if "table" in body:
        table = body["table"]
        rows = table["rows"] + ([table["total"]] if "total" in table else [])
        # A leading "No." column is just numbering; the next column names the row.
        key = 1 if table["columns"][0]["label"] == "No." and len(table["columns"]) > 2 else 0
        for row in rows:
            category = row[key] or row[0]
            for col, value in zip(table["columns"][key + 1:], row[key + 1:]):
                yield col["label"], category, value
        return
    chart = body["chart"]
    if "values" in chart:
        for label, value in zip(chart["labels"], chart["values"]):
            yield default_series, label, value
        return
    for ds in chart["datasets"]:
        for label, value in zip(chart["labels"], ds["data"]):
            yield ds["label"], label, value


def long_rows(fig):
    """One row per value across the figure; a panel's title prefixes its series."""
    for title, body in bodies(fig):
        for series, category, value in body_long_rows(body, title or fig["caption"]):
            if title and series != title:
                series = f"{title}: {series}"
            yield series, category, value


def sheet_name(fid):
    return f"Fig {fid}"


def autosize(ws, min_width=8, max_width=60):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_width, min(width + 2, max_width))


def write_header(ws, row, values):
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = BOLD
        cell.fill = HEADER_FILL


def build_xlsx(master):
    figures = master["figures"]
    wb = Workbook()

    index = wb.active
    index.title = "Index"
    index["A1"] = "IPR 2025 figures"
    index["A1"].font = TITLE
    index["A2"] = master["source"]
    for i, note in enumerate(master["notes"]):
        index.cell(row=3 + i, column=1, value=note)
    top = 4 + len(master["notes"])
    columns = ["ID", "Chapter", "Chapter name", "Caption", "Page", "Type", "Source", "Note", "Estimated", "Flags", "Sheet"]
    write_header(index, top, columns)

    for r, fig in enumerate(figures, start=top + 1):
        has_data = "skipped" not in fig
        values = [
            fig["id"], fig["chapter"], fig["chapterName"], fig["caption"], fig.get("page"),
            figure_kind(fig), fig.get("source"), fig.get("note") or fig.get("skipped"),
            "yes" if fig.get("estimated") else "", " | ".join(fig.get("flags", [])),
            sheet_name(fig["id"]) if has_data else "",
        ]
        for c, value in enumerate(values, start=1):
            index.cell(row=r, column=c, value=value)
        if has_data:
            link = index.cell(row=r, column=len(columns))
            link.hyperlink = f"#'{sheet_name(fig['id'])}'!A1"
            link.style = "Hyperlink"
        if fig.get("flags"):
            index.cell(row=r, column=10).font = FLAG_FONT

    autosize(index)
    for col, width in (("D", 60), ("H", 60), ("J", 60)):
        index.column_dimensions[col].width = width
    for row in index.iter_rows(min_row=top + 1):
        for cell in row:
            cell.alignment = WRAP
    index.freeze_panes = index.cell(row=top + 1, column=2)

    for fig in figures:
        if "skipped" in fig:
            continue
        ws = wb.create_sheet(sheet_name(fig["id"]))
        ws["A1"] = f"Figure {fig['id']}: {fig['caption']}"
        ws["A1"].font = TITLE
        meta = [("Chapter", f"{fig['chapter']} – {fig['chapterName']}"), ("Page", fig.get("page")), ("Source", fig.get("source"))]
        if fig.get("note"):
            meta.append(("Note", fig["note"]))
        if fig.get("estimated"):
            meta.append(("Estimated", "Values measured from the chart; the report prints no data labels."))
        for flag in fig.get("flags", []):
            meta.append(("Flag", flag))
        for i, (key, value) in enumerate(meta, start=2):
            ws.cell(row=i, column=1, value=key).font = BOLD
            cell = ws.cell(row=i, column=2, value=value)
            if key == "Flag":
                cell.font = FLAG_FONT

        # One block per chart or table; panels each get their title above.
        r = len(meta) + 3
        first_header = None
        for title, body in bodies(fig):
            if title:
                ws.cell(row=r, column=1, value=title).font = BOLD
                r += 1
            header, rows, total = body_rows(body)
            write_header(ws, r, header)
            first_header = first_header or r
            for row in rows + ([total] if total else []):
                r += 1
                for c, value in enumerate(row, start=1):
                    cell = ws.cell(row=r, column=c, value=value)
                    if row is total:
                        cell.font = BOLD
            r += 2
        autosize(ws)
        ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width, 40)
        if "panels" not in fig:
            ws.freeze_panes = ws.cell(row=first_header + 1, column=2)

    long = wb.create_sheet("All data (long)")
    write_header(long, 1, ["Figure ID", "Caption", "Series", "Category", "Value", "Estimated"])
    r = 2
    for fig in figures:
        if "skipped" in fig:
            continue
        for series, category, value in long_rows(fig):
            long.append([fig["id"], fig["caption"], series, category, value, "yes" if fig.get("estimated") else ""])
            r += 1
    autosize(long)
    long.freeze_panes = "A2"
    long.auto_filter.ref = f"A1:F{r - 1}"

    flags = wb.create_sheet("Flags")
    write_header(flags, 1, ["Figure ID", "Caption", "Page", "Issue"])
    for fig in figures:
        issues = list(fig.get("flags", []))
        if fig.get("estimated"):
            issues.insert(0, "Estimated: values measured from the chart (no data labels in the report).")
        for issue in issues:
            flags.append([fig["id"], fig["caption"], fig.get("page"), issue])
    autosize(flags)
    flags.column_dimensions["D"].width = 100
    for row in flags.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP

    wb.save(XLSX)


def js_value(value, indent=0):
    """Serialise to JS literal syntax matching figures.js (single quotes, bare keys)."""
    pad = "  " * indent
    inner = "  " * (indent + 1)
    if value is None:
        return "''"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return repr(value)
    if isinstance(value, str):
        return "'" + value.replace("\\", "\\\\").replace("'", "\\'") + "'"
    if isinstance(value, list):
        if all(not isinstance(v, (dict, list)) for v in value):
            return "[" + ", ".join(js_value(v) for v in value) + "]"
        return "[\n" + ",\n".join(inner + js_value(v, indent + 1) for v in value) + ",\n" + pad + "]"
    if isinstance(value, dict):
        def key(k):
            return k if re.fullmatch(r"[A-Za-z_$][\w$]*", k) else js_value(k)
        items = [f"{inner}{key(k)}: {js_value(v, indent + 1)}" for k, v in value.items()]
        return "{\n" + ",\n".join(items) + ",\n" + pad + "}"
    raise TypeError(type(value))


SITE_KEYS = (
    "id", "chapter", "chapterName", "caption", "source", "note", "chart", "table",
    "sharedLegend", "panelColumns", "panelMinWidth", "panelRatio", "panels",
)


def build_js(master):
    figures = [
        {k: fig[k] for k in SITE_KEYS if k in fig}
        for fig in master["figures"]
        if "skipped" not in fig
    ]
    header = (
        "// Generated by data-extraction/build.py from data-extraction/figures-data.json.\n"
        "// Don't edit by hand: change figures-data.json and rerun the build.\n"
        "// Photos and diagrams without data are left out; see the workbook's Index sheet.\n"
    )
    # Series colours for datasets that don't set their own, per chapter, from the report.
    palettes = {c: chapter["palette"] for c, chapter in master["chapters"].items()}
    JS.write_text(
        header
        + "export const chapterPalettes = " + js_value(palettes) + "\n\n"
        + "export const figures = " + js_value(figures) + "\n",
        encoding="utf-8",
    )
    return len(figures)


def main():
    master = json.loads(DATA.read_text(encoding="utf-8"))
    validate(master["figures"])
    build_xlsx(master)
    count = build_js(master)
    print(f"Wrote {XLSX.relative_to(ROOT)} and {JS.relative_to(ROOT)} ({count} figures with data).")


if __name__ == "__main__":
    main()
